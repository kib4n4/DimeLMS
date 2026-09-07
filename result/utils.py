from openpyxl import Workbook, load_workbook

from .models import TakenCourse, Result
from core.models import Session, Semester
from course.models import Course

SCORE_BULK_UPLOAD_COLUMNS = [
    "student id",
    "assignment",
    "mid exam",
    "quiz",
    "attendance",
    "final exam",
]

# Score fields are entered as raw marks out of 100 for each component, the
# same as the single-student "Manage Score" form's number inputs.
_SCORE_FIELDS = ["assignment", "mid_exam", "quiz", "attendance", "final_exam"]
_SCORE_LABELS = {
    "assignment": "Assignment",
    "mid_exam": "Mid Exam",
    "quiz": "Quiz",
    "attendance": "Attendance",
    "final_exam": "Final Exam",
}


def build_score_bulk_upload_template(course):
    """
    An .xlsx workbook pre-filled with one row per student currently taking
    `course` in the active semester — their existing scores included, so a
    facilitator can download it, edit marks in a spreadsheet, and re-upload
    it rather than typing every field into the on-page table.
    """
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Scores"
    sheet.append(
        ["Student ID", "Student Name", "Assignment", "Mid Exam", "Quiz", "Attendance", "Final Exam"]
    )

    current_semester = Semester.objects.filter(is_current_semester=True).first()
    taken_courses = TakenCourse.objects.filter(course=course)
    if current_semester:
        taken_courses = taken_courses.filter(course__semester=current_semester)

    for taken_course in taken_courses.select_related("student__student"):
        student_user = taken_course.student.student
        sheet.append(
            [
                student_user.username,
                student_user.get_full_name,
                float(taken_course.assignment),
                float(taken_course.mid_exam),
                float(taken_course.quiz),
                float(taken_course.attendance),
                float(taken_course.final_exam),
            ]
        )

    if sheet.max_row == 1:
        # No students registered for the course yet — leave one example row
        # so the column layout is still obvious.
        sheet.append(["student-id-here", "", 0, 0, 0, 0, 0])

    for column_cells in sheet.columns:
        length = max(len(str(cell.value)) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = length + 4

    return wb


def parse_score_bulk_upload(uploaded_file, course, lecturer):
    """
    Parse an uploaded .xlsx file of scores (columns: Student ID, Assignment,
    Mid Exam, Quiz, Attendance, Final Exam) for `course`, matching each row
    to an existing TakenCourse record by the student's username — a student
    must already be registered for the course (i.e. have a TakenCourse row)
    before a score can be recorded for them, same as the single-entry form.

    Saves valid rows immediately, computing total/grade/point/comment and
    the student's GPA/CGPA the same way `add_score_for`'s POST handler
    does. Returns (updated_count, row_errors) — row_errors is a list of
    {"row": int, "message": str} dicts, one per row that couldn't be saved.
    """
    try:
        wb = load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception:
        return 0, [{"row": 1, "message": "Couldn't read that file — make sure it's a valid .xlsx workbook."}]

    sheet = wb.active
    rows = sheet.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        return 0, [{"row": 1, "message": "The file is empty."}]

    headers = [str(cell).strip().lower() if cell is not None else "" for cell in header_row]
    column_index = {name: headers.index(name) for name in SCORE_BULK_UPLOAD_COLUMNS if name in headers}

    missing_columns = [name.title() for name in SCORE_BULK_UPLOAD_COLUMNS if name not in column_index]
    if missing_columns:
        return 0, [{
            "row": 1,
            "message": f"Missing required column(s): {', '.join(missing_columns)}.",
        }]

    def cell_value(row, field):
        index = column_index[field]
        return row[index] if index < len(row) else None

    current_session = Session.objects.filter(is_current_session=True).first()
    current_semester = Semester.objects.filter(
        is_current_semester=True, session=current_session
    ).first()

    updated_count = 0
    row_errors = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue  # skip fully blank rows

        student_id = cell_value(row, "student id")
        student_id = str(student_id).strip() if student_id is not None else ""
        if not student_id:
            row_errors.append({"row": row_number, "message": "Missing Student ID."})
            continue

        taken_course = (
            TakenCourse.objects.filter(
                course=course,
                course__allocated_course__lecturer=lecturer,
                student__student__username=student_id,
            )
            .select_related("student", "student__student", "course")
            .first()
        )
        if not taken_course:
            row_errors.append({
                "row": row_number,
                "message": f'"{student_id}" isn\'t registered for this course.',
            })
            continue

        scores = {}
        invalid_field = None
        for field in _SCORE_FIELDS:
            value = cell_value(row, field.replace("_", " "))
            if value is None or value == "":
                value = 0
            try:
                value = float(value)
            except (TypeError, ValueError):
                invalid_field = field
                break
            if value < 0 or value > 100:
                invalid_field = field
                break
            scores[field] = value

        if invalid_field:
            row_errors.append({
                "row": row_number,
                "message": f"{_SCORE_LABELS[invalid_field]} must be a number between 0 and 100.",
            })
            continue

        taken_course.assignment = scores["assignment"]
        taken_course.mid_exam = scores["mid_exam"]
        taken_course.quiz = scores["quiz"]
        taken_course.attendance = scores["attendance"]
        taken_course.final_exam = scores["final_exam"]
        taken_course.total = taken_course.get_total(**scores)
        taken_course.grade = taken_course.get_grade(total=taken_course.total)
        taken_course.point = taken_course.get_point(grade=taken_course.grade)
        taken_course.comment = taken_course.get_comment(grade=taken_course.grade)
        taken_course.save()
        updated_count += 1

        if current_semester and current_session:
            student_profile = taken_course.student
            courses_in_semester = Course.objects.filter(
                level=student_profile.level,
                program__pk=student_profile.program_id,
                semester=current_semester,
            )
            total_credit_in_semester = sum(int(c.credit) for c in courses_in_semester)
            gpa = taken_course.calculate_gpa(total_credit_in_semester)
            cgpa = taken_course.calculate_cgpa()
            try:
                result = Result.objects.get(
                    student=student_profile,
                    semester=current_semester,
                    session=current_session,
                    level=student_profile.level,
                )
                result.gpa = gpa
                result.cgpa = cgpa
                result.save()
            except Result.DoesNotExist:
                Result.objects.get_or_create(
                    student=student_profile,
                    gpa=gpa,
                    semester=current_semester,
                    session=current_session,
                    level=student_profile.level,
                )

    return updated_count, row_errors
