import datetime
import os
import random
import re
import string

import docx
from django.core.exceptions import ValidationError
from django.utils.text import slugify

YOUTUBE_URL_PATTERN = re.compile(
    r"(?:youtube\.com/(?:watch\?v=|embed/|v/|shorts/)|youtu\.be/)([A-Za-z0-9_-]{11})"
)


def extract_youtube_id(url):
    """Pull the 11-char video id out of any common YouTube URL shape
    (watch?v=, youtu.be/, embed/, shorts/), or None if it doesn't match."""
    if not url:
        return None
    match = YOUTUBE_URL_PATTERN.search(url)
    return match.group(1) if match else None


def validate_youtube_url(value):
    if not extract_youtube_id(value):
        raise ValidationError(
            "Enter a valid YouTube link, e.g. https://www.youtube.com/watch?v=... or https://youtu.be/..."
        )


def random_string_generator(size=10, chars=string.ascii_lowercase + string.digits):
    return "".join(random.choice(chars) for _ in range(size))


def unique_slug_generator(instance, new_slug=None):
    """
    This is for a Django project and it assumes your instance
    has a model with a slug field and a title character (char) field.
    """
    if new_slug is not None:
        slug = new_slug
    else:
        slug = slugify(instance.title)

    Klass = instance.__class__
    qs_exists = Klass.objects.filter(slug=slug).exists()
    if qs_exists:
        new_slug = "{slug}-{randstr}".format(
            slug=slug, randstr=random_string_generator(size=4)
        )
        return unique_slug_generator(instance, new_slug=new_slug)
    return slug


def extract_docx_paragraphs(file_obj):
    """Returns the paragraph text of a .docx file, for the in-page document
    reader (see Upload.viewer_kind / course.views.document_single) — plain
    text only, no formatting/images, since there's no public URL for
    Google/Office's online viewers to render the original from on a local
    dev server."""
    document = docx.Document(file_obj)
    return [p.text for p in document.paragraphs if p.text.strip()]


def split_docx_into_topics(file_obj):
    """
    Splits a .docx into topics using its Word "Heading 1" style — each
    Heading-1 paragraph starts a new topic, and everything until the next
    Heading 1 belongs to it, sub-headings (Heading 2, 3, ...) included as
    part of that topic's body rather than as topics of their own. Used to
    auto-generate one course Module per top-level topic (see
    course.views.module_split_view).

    Only Heading 1 is a boundary — not "Title" or lower heading levels —
    so a document with a non-numbered title block and numbered sections
    that themselves contain lettered/numbered sub-headings (e.g. "4.1",
    "4.2") produces one card per top-level section, not one per
    sub-heading, and the title block is dropped as preamble.

    Returns a list of {"title": str, "content": str} dicts, one per
    detected topic, in document order. A document with no Heading-1
    paragraphs at all returns an empty list; the caller is responsible
    for telling the facilitator nothing was found.
    """
    document = docx.Document(file_obj)
    topics = []
    current = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = (paragraph.style.name or "").strip().lower()

        if style_name == "heading 1":
            if current is not None:
                topics.append(current)
            current = {"title": text, "content": ""}
            continue

        if current is None:
            continue  # ignore any preamble before the first Heading 1

        current["content"] += text + "\n\n"

    if current is not None:
        topics.append(current)

    for topic in topics:
        topic["content"] = topic["content"].strip()

    return topics
